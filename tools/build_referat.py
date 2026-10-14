# -*- coding: utf-8 -*-
"""Сборщик реферата в формате шаблонов НовГУ (ТNR 14, 1.5, поля 3/1.5/2/2)."""
import math, os, sys
from PIL import Image
from docx import Document
from docx.shared import Pt, Cm, Twips, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from referat_content_1 import BLOCKS_1
from referat_content_2 import BLOCKS_2
from referat_content_3 import BLOCKS_3

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG = os.path.join(ROOT, "ОВЧ-ЭКЗ", "Итоговые работы", "images")
OUT = os.path.join(ROOT, "ОВЧ-ЭКЗ", "Итоговые работы", "3952_Теддер_Д_ОхлаждениеКомпьютерныхСистем_Реферат.docx")

BLOCKS = BLOCKS_1 + BLOCKS_2 + BLOCKS_3

LINE_CM = 14 * 1.5 * 0.0352778          # высота строки текста, см
LINES_PAGE = 34
CHARS_LINE = 64
TBL_LINE_CM = 12 * 1.15 * 0.0352778     # строка в таблице 12 пт
CHARS_TBL = 0.215                        # ширина символа 12 пт, см

# ---------------------------------------------------------------- пагинация
def img_size_cm(key, width_cm):
    for ext in ("jpg", "png"):
        p = os.path.join(IMG, f"{key}.{ext}")
        if os.path.exists(p):
            w, h = Image.open(p).size
            return width_cm, width_cm * h / w
    return width_cm, width_cm

def para_lines(text, chars=CHARS_LINE):
    return max(1, math.ceil(len(text) / chars))

def paginate(blocks):
    """Возвращает (pages: dict key->страница, total)."""
    page, line = 2, 0   # титул + страница содержания
    marks = {}
    for b in blocks:
        kind = b[0]
        if kind in ("h1_c", "h1"):
            page += 1; line = 0
            marks[b[1]] = page
            line += para_lines(b[1])
        elif kind == "h2":
            if line + 2 > LINES_PAGE:
                page += 1; line = 0
            marks[b[1]] = page
            line += para_lines(b[1])
        elif kind == "p":
            n = para_lines(b[1])
            if line + n > LINES_PAGE:
                page += 1; line = 0
            line += n
        elif kind == "tbl":
            caption, headers, rows = b[1], b[2], b[3]
            ncols = len(headers)
            allrows = [headers] + rows
            avg = [sum(len(r[c]) for r in allrows) / len(allrows) for c in range(ncols)]
            tot = sum(avg) or 1
            colw = [16.5 * a / tot for a in avg]
            colw = [max(c, 2.2) for c in colw]
            def row_lines(r):
                return max(max(1, math.ceil(len(c) / max(6, (colw[i] - 0.35) / CHARS_TBL)))
                           for i, c in enumerate(r))
            h_cm = sum(row_lines(r) * TBL_LINE_CM + 0.15 for r in allrows)
            need = h_cm / LINE_CM + para_lines(caption)
            if line + need > LINES_PAGE:
                page += 1; line = 0
            line += need
        elif kind == "fig":
            key, caption, wcm = b[1], b[2], b[3]
            _, hcm = img_size_cm(key, wcm)
            need = (hcm + 0.4) / LINE_CM + para_lines(caption)
            if line + need > LINES_PAGE:
                page += 1; line = 0
            line += need
    return marks, page

MARKS, TOTAL = paginate(BLOCKS)
print("Оценка объёма: страниц =", TOTAL)

# ---------------------------------------------------------------- документ
doc = Document()

# стили
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(14)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
npf = normal.paragraph_format
npf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
npf.line_spacing = 1.5
npf.space_before = Pt(0)
npf.space_after = Pt(0)

for sname in ("Title", "Subtitle"):
    st = doc.styles[sname]
    st.font.name = "Times New Roman"
    st.font.size = Pt(14)
    st.font.bold = False
    st.font.color.rgb = RGBColor(0, 0, 0)
    spf = st.paragraph_format
    spf.line_spacing = 1.5
    spf.space_before = Pt(0)
    spf.space_after = Pt(0)
    spf.keep_with_next = True

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin, sec.right_margin = Cm(3.0), Cm(1.5)
sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)

# колонтитул: номер страницы сверху по центру
hdr = sec.header
hp = hdr.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
hp.paragraph_format.line_spacing = 1.0
r = hp.add_run()
fb = OxmlElement("w:fldChar"); fb.set(qn("w:fldCharType"), "begin"); r._r.append(fb)
r2 = hp.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
it.text = " PAGE   \\* MERGEFORMAT "; r2._r.append(it)
r3 = hp.add_run(); fs = OxmlElement("w:fldChar"); fs.set(qn("w:fldCharType"), "separate"); r3._r.append(fs)
r4 = hp.add_run("1")
r5 = hp.add_run(); fe = OxmlElement("w:fldChar"); fe.set(qn("w:fldCharType"), "end"); r5._r.append(fe)
for rr in (r, r2, r3, r4, r5):
    rr.font.name = "Times New Roman"; rr.font.size = Pt(14)

# обновление полей при открытии
settings = doc.settings.element
uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
settings.append(uf)

def set_first_line(p, tw):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind"); pPr.append(ind)
    ind.set(qn("w:firstLine"), str(tw))

def add_para(text, style=None, align=None, first_tw=None, page_break=False, line=1.5):
    p = doc.add_paragraph(style=style)
    if text:
        p.add_run(text)
    pf = p.paragraph_format
    if align is not None:
        pf.alignment = align
    pf.line_spacing = line
    if first_tw is not None:
        set_first_line(p, first_tw)
    if page_break:
        pf.page_break_before = True
    return p

# ---------------------------------------------------------------- титульный лист
title_lines = [
    "Министерство науки и высшего образования Российской Федерации",
    "Федеральное государственное бюджетное образовательное учреждение",
    "высшего образования",
    "«Новгородский государственный университет имени Ярослава Мудрого»",
    "ПОЛИТЕХНИЧЕСКИЙ ИНСТИТУТ",
    "ПОЛИТЕХНИЧЕСКИЙ КОЛЛЕДЖ",
]
for t in title_lines:
    add_para(t, align=WD_ALIGN_PARAGRAPH.CENTER, first_tw=0, line=1.1)
add_para("", line=1.1)
p = add_para("РЕФЕРАТ", align=WD_ALIGN_PARAGRAPH.CENTER, first_tw=0, line=1.1)
p.runs[0].font.size = Pt(16)
p = add_para("Анализ технологий охлаждения компьютерных систем и основы их обслуживания",
             align=WD_ALIGN_PARAGRAPH.CENTER, first_tw=0, line=1.1)
p.runs[0].font.size = Pt(16)
add_para("по дисциплине «Техническое обслуживание и ремонт аппаратной части компьютерных систем и комплексов»",
         align=WD_ALIGN_PARAGRAPH.CENTER, first_tw=0, line=1.1)
add_para("По специальности 09.02.01 Компьютерные системы и комплексы",
         align=WD_ALIGN_PARAGRAPH.CENTER, first_tw=0, line=1.1)
for _ in range(3):
    add_para("", line=1.1)
for t in ["Руководитель", "Овчинников К.Е. ____________", "(подпись)", "«___» ____________ 20___ года",
          "Выполнил: обучающийся", "группы №3952", "Теддер Д.___________", "(подпись)", "«___» ____________ 20___ года"]:
    add_para(t, align=WD_ALIGN_PARAGRAPH.RIGHT, first_tw=0, line=1.1)
for _ in range(3):
    add_para("", line=1.1)
add_para("Великий Новгород, 2026", align=WD_ALIGN_PARAGRAPH.CENTER, first_tw=0, line=1.1)

# ---------------------------------------------------------------- содержание
add_para("Содержание", align=WD_ALIGN_PARAGRAPH.CENTER, first_tw=0, page_break=True)

toc_entries = []
for b in BLOCKS:
    if b[0] in ("h1_c", "h1", "h2"):
        toc_entries.append((b[1], MARKS[b[1]], b[0] == "h2"))

# поле TOC с кэшированными записями
p = doc.add_paragraph()
pf = p.paragraph_format; pf.line_spacing = 1.5
r = p.add_run(); fb = OxmlElement("w:fldChar"); fb.set(qn("w:fldCharType"), "begin"); r._r.append(fb)
r = p.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
it.text = ' TOC \\o "1-3" \\h \\z \\u '; r._r.append(it)
r = p.add_run(); fs = OxmlElement("w:fldChar"); fs.set(qn("w:fldCharType"), "separate"); r._r.append(fs)

def toc_line(text, page, level2):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pPr = p._p.get_or_add_pPr()
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "0"); ind.set(qn("w:firstLine"), "0")
    pPr.append(ind)
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right"); tab.set(qn("w:pos"), "9355"); tab.set(qn("w:leader"), "dot")
    tabs.append(tab); pPr.append(tabs)
    r = p.add_run(("\t" if False else "") + text)
    if level2:
        r.font.size = Pt(14)
    r2 = p.add_run("\t" + str(page))
    return p

for text, page, level2 in toc_entries:
    toc_line(text, page, level2)
p = doc.add_paragraph()
r = p.add_run(); fe = OxmlElement("w:fldChar"); fe.set(qn("w:fldCharType"), "end"); r._r.append(fe)

# ---------------------------------------------------------------- тело
def add_table(caption, headers, rows):
    add_para(caption, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_tw=0)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    tbl = t._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        borders.append(e)
    tblPr.append(borders)
    hdr_tr = t.rows[0]._tr
    trPr = hdr_tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); trPr.append(th)
    def fill(cell, text, center=False):
        cell.paragraphs[0].text = ""
        p = cell.paragraphs[0]
        p.add_run(text)
        pf = p.paragraph_format
        pf.line_spacing = 1.0
        pf.space_before = Pt(0); pf.space_after = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        set_first_line(p, 0)
        for run in p.runs:
            run.font.size = Pt(12)
            run.font.name = "Times New Roman"
    for i, h in enumerate(headers):
        fill(t.rows[0].cells[i], h, center=True)
    for row in rows:
        cells = t.add_row().cells
        for i, c in enumerate(row):
            fill(cells[i], c, center=(i > 0))
    return t

def add_figure(key, caption, width_cm):
    path = None
    for ext in ("jpg", "png"):
        pth = os.path.join(IMG, f"{key}.{ext}")
        if os.path.exists(pth):
            path = pth
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    set_first_line(p, 0)
    p.add_run().add_picture(path, width=Cm(width_cm))
    add_para(caption, align=WD_ALIGN_PARAGRAPH.CENTER, first_tw=0)

for b in BLOCKS:
    kind = b[0]
    if kind == "h1_c":
        add_para(b[1], style="Title", align=WD_ALIGN_PARAGRAPH.CENTER, first_tw=0, page_break=True)
    elif kind == "h1":
        add_para(b[1], style="Title", align=WD_ALIGN_PARAGRAPH.LEFT, first_tw=720, page_break=True)
    elif kind == "h2":
        add_para(b[1], style="Subtitle", align=WD_ALIGN_PARAGRAPH.LEFT, first_tw=720)
    elif kind == "p":
        add_para(b[1], first_tw=709)
    elif kind == "tbl":
        add_table(b[1], b[2], b[3])
    elif kind == "fig":
        add_figure(b[1], b[2], b[3])

doc.core_properties.title = "Анализ технологий охлаждения компьютерных систем и основы их обслуживания"
doc.core_properties.author = "Теддер Даниил"
doc.save(OUT)
print("Сохранено:", OUT)
print("Записей в содержании:", len(toc_entries))
